"""
xlsx_parser.py - Excel parser.

Strategy:
1. Try openpyxl with data_only=True.
2. If that yields no sheet rows, try openpyxl with data_only=False.
3. If the workbook is still unreadable (for example strict OOXML files that
   the installed openpyxl version fails to materialize into worksheets),
   fall back to reading workbook.xml / sharedStrings.xml / sheet XML directly.
"""

from __future__ import annotations

import logging
import posixpath
import warnings
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from app.schemas import Chunk, ParsedDocument, normalize_text, stable_id


logger = logging.getLogger(__name__)

DOC_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
STRICT_DOC_REL_NS = "http://purl.oclc.org/ooxml/officeDocument/relationships"


def _build_doc_id(path: Path) -> str:
    stat = path.stat()
    return stable_id(path.resolve(), stat.st_size, int(stat.st_mtime))


def _detect_header_row(rows: list[tuple[object, ...]]) -> tuple[int, list[str]]:
    for index, row in enumerate(rows):
        values = [normalize_text(cell) for cell in row]
        non_empty = [value for value in values if value]
        if len(non_empty) >= 2:
            return index, [value or f"列{column_index + 1}" for column_index, value in enumerate(values)]
    return 0, [f"列{column_index + 1}" for column_index in range(len(rows[0]) if rows else 0)]


def _maybe_format_cell(header: str, value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, (int, float)) and any(token in header.lower() for token in ["date", "time", "日期", "时间"]):
        try:
            return from_excel(value).strftime("%Y-%m-%d %H:%M:%S")
        except Exception:
            return normalize_text(value)
    return normalize_text(value)


def _safe_load_workbook(path: str, data_only: bool):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        return load_workbook(path, data_only=data_only)


def _parse_xlsx_with_workbook(target: Path, doc_id: str, workbook, parser_backend: str) -> tuple[list[Chunk], list[str]]:
    chunks: list[Chunk] = []
    full_text: list[str] = []

    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue

        header_row_index, headers = _detect_header_row(rows)
        for row_index, row in enumerate(rows[header_row_index + 1 :], start=header_row_index + 2):
            structured_row = {}
            for column_index, cell in enumerate(row):
                header = headers[column_index] if column_index < len(headers) else f"列{column_index + 1}"
                value = _maybe_format_cell(header, cell)
                if value:
                    structured_row[header] = value

            if not structured_row:
                continue

            row_text = f"Sheet={sheet.title} | " + " | ".join(
                f"{key}: {value}"
                for key, value in structured_row.items()
            )
            full_text.append(row_text)
            chunks.append(
                Chunk(
                    chunk_id=f"{doc_id}:s:{sheet.title}:r:{row_index}",
                    doc_id=doc_id,
                    text=row_text,
                    metadata={
                        "file_path": str(target.resolve()),
                        "file_type": "xlsx",
                        "block_type": "sheet_row",
                        "sheet_name": sheet.title,
                        "row_index": row_index,
                        "headers": headers,
                        "structured_row": structured_row,
                        "parser_backend": parser_backend,
                    },
                )
            )

    return chunks, full_text


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _column_index_from_ref(cell_ref: str) -> int:
    letters = "".join(char for char in cell_ref if char.isalpha()).upper()
    index = 0
    for char in letters:
        index = index * 26 + (ord(char) - ord("A") + 1)
    return max(index - 1, 0)


def _extract_shared_string(si_element: ET.Element) -> str:
    texts: list[str] = []
    for node in si_element.iter():
        if _local_name(node.tag) == "t" and node.text:
            texts.append(node.text)
    return "".join(texts).strip()


def _parse_archive_shared_strings(archive: ZipFile) -> list[str]:
    try:
        root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
    except KeyError:
        return []
    return [_extract_shared_string(si) for si in root if _local_name(si.tag) == "si"]


def _parse_archive_rows(sheet_xml: bytes, shared_strings: list[str]) -> list[tuple[object, ...]]:
    root = ET.fromstring(sheet_xml)
    rows: list[tuple[object, ...]] = []

    for row_element in root.iter():
        if _local_name(row_element.tag) != "row":
            continue

        row_values: dict[int, object] = {}
        max_column = -1
        for cell in row_element:
            if _local_name(cell.tag) != "c":
                continue

            cell_ref = cell.attrib.get("r", "")
            column_index = _column_index_from_ref(cell_ref) if cell_ref else max_column + 1
            max_column = max(max_column, column_index)
            cell_type = cell.attrib.get("t", "")
            value_text = ""
            inline_text = ""

            for child in cell:
                tag = _local_name(child.tag)
                if tag == "v":
                    value_text = child.text or ""
                elif tag == "is":
                    inline_text = "".join(
                        node.text or ""
                        for node in child.iter()
                        if _local_name(node.tag) == "t"
                    )

            if cell_type == "s":
                try:
                    value: object = shared_strings[int(value_text or "0")]
                except Exception:
                    value = value_text
            elif cell_type == "inlineStr":
                value = inline_text
            elif cell_type == "b":
                value = "TRUE" if value_text == "1" else "FALSE"
            else:
                value = value_text or inline_text

            row_values[column_index] = value

        if max_column < 0:
            continue

        rows.append(tuple(row_values.get(column_index, "") for column_index in range(max_column + 1)))

    return rows


def _resolve_sheet_targets(archive: ZipFile) -> list[tuple[str, str]]:
    workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
    rels_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))

    rel_map: dict[str, str] = {}
    for rel in rels_root:
        if _local_name(rel.tag) != "Relationship":
            continue
        rel_id = rel.attrib.get("Id", "")
        target = rel.attrib.get("Target", "")
        rel_type = rel.attrib.get("Type", "")
        if rel_id and target and "worksheet" in rel_type:
            rel_map[rel_id] = target

    sheet_targets: list[tuple[str, str]] = []
    for element in workbook_root.iter():
        if _local_name(element.tag) != "sheet":
            continue
        sheet_name = element.attrib.get("name", "Sheet")
        rel_id = (
            element.attrib.get(f"{{{DOC_REL_NS}}}id")
            or element.attrib.get(f"{{{STRICT_DOC_REL_NS}}}id")
            or element.attrib.get("id")
            or element.attrib.get("r:id")
        )
        target = rel_map.get(rel_id or "")
        if not target:
            continue

        normalized_target = target.lstrip("/")
        if not normalized_target.startswith("xl/"):
            normalized_target = posixpath.normpath(posixpath.join("xl", normalized_target))
        sheet_targets.append((sheet_name, normalized_target))

    return sheet_targets


def _parse_xlsx_from_archive(target: Path, doc_id: str) -> tuple[list[Chunk], list[str]]:
    chunks: list[Chunk] = []
    full_text: list[str] = []

    with ZipFile(target) as archive:
        shared_strings = _parse_archive_shared_strings(archive)
        for sheet_name, sheet_target in _resolve_sheet_targets(archive):
            rows = _parse_archive_rows(archive.read(sheet_target), shared_strings)
            if not rows:
                continue

            header_row_index, headers = _detect_header_row(rows)
            for row_index, row in enumerate(rows[header_row_index + 1 :], start=header_row_index + 2):
                structured_row = {}
                for column_index, cell in enumerate(row):
                    header = headers[column_index] if column_index < len(headers) else f"列{column_index + 1}"
                    value = _maybe_format_cell(header, cell)
                    if value:
                        structured_row[header] = value

                if not structured_row:
                    continue

                row_text = f"Sheet={sheet_name} | " + " | ".join(
                    f"{key}: {value}"
                    for key, value in structured_row.items()
                )
                full_text.append(row_text)
                chunks.append(
                    Chunk(
                        chunk_id=f"{doc_id}:s:{sheet_name}:r:{row_index}",
                        doc_id=doc_id,
                        text=row_text,
                        metadata={
                            "file_path": str(target.resolve()),
                            "file_type": "xlsx",
                            "block_type": "sheet_row",
                            "sheet_name": sheet_name,
                            "row_index": row_index,
                            "headers": headers,
                            "structured_row": structured_row,
                            "parser_backend": "zip_xml",
                        },
                    )
                )

    return chunks, full_text


def parse_xlsx(path: str) -> ParsedDocument:
    target = Path(path)
    doc_id = _build_doc_id(target)

    chunks: list[Chunk] = []
    full_text: list[str] = []

    try:
        workbook = _safe_load_workbook(path, data_only=True)
        chunks, full_text = _parse_xlsx_with_workbook(target, doc_id, workbook, parser_backend="openpyxl_data_only")
    except Exception as exc:
        logger.warning("XLSX data_only 解析失败，将尝试回退: %s (%s)", target.name, exc)

    if not full_text:
        try:
            workbook2 = _safe_load_workbook(path, data_only=False)
            chunks, full_text = _parse_xlsx_with_workbook(target, doc_id, workbook2, parser_backend="openpyxl_formula")
        except Exception as exc:
            logger.warning("XLSX formula 解析失败，将尝试 XML 回退: %s (%s)", target.name, exc)

    if not full_text:
        logger.info("XLSX workbook 无可见工作表或 openpyxl 未读出数据，使用 ZIP/XML 回退解析: %s", target.name)
        chunks, full_text = _parse_xlsx_from_archive(target, doc_id)

    return ParsedDocument(
        doc_id=doc_id,
        name=target.name,
        path=str(target),
        file_type="xlsx",
        text="\n".join(full_text),
        chunks=chunks,
        metadata={"last_modified": target.stat().st_mtime},
    )
