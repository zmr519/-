"""
filler/utils.py — 模板解析与填写工具

提供：
  - inspect_template()     : 解析模板结构（占位符字段 + 表格表头 + 目标城市）
  - get_template_fields()  : 获取模板中所有字段名
  - build_output_path()    : 生成输出文件路径
  - save_json()            : 保存审计报告 JSON
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
import json
import re

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from openpyxl import load_workbook

from app.city_utils import detect_target_city
from app.schemas import TemplateLayout, TemplateTableTarget, normalize_field_name, normalize_text


# 占位符正则：匹配 {{field_name}} 样式
PLACEHOLDER_PATTERN = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")

# 中文括号占位符：匹配【field_description】样式
CN_PLACEHOLDER_PATTERN = re.compile(r"【([^【】]+?)】")

# 下划线填空占位符：匹配 2 个以上连续下划线
UNDERLINE_BLANK_PATTERN = re.compile(r"_{2,}")

# 方括号填空占位符：匹配 [ ] 或 [ ] 等空内容方括号
BRACKET_BLANK_PATTERN = re.compile(r"\[\s*\]")

# 中文圆括号填空占位符：匹配（）内容为空或仅含下划线/空格
CN_PAREN_BLANK_PATTERN = re.compile(r"（\s*）|（_{2,}）|（\s+）")

# 表格中的通用填写标记（不含字段描述）
_GENERIC_TABLE_MARKERS = {"填写", "合计填写", "合计", "合计变化"}


def inspect_template(template_path: str) -> TemplateLayout:
    """解析模板文件，返回 TemplateLayout（包含占位符字段、表格表头等）。"""
    path = Path(template_path)
    suffix = path.suffix.lower()
    if suffix == ".xlsx":
        return _inspect_excel_template(path)
    if suffix == ".docx":
        return _inspect_word_template(path)
    raise ValueError(f"不支持的模板类型: {template_path}")


def get_template_fields(template_path: str) -> list[str]:
    return inspect_template(template_path).all_fields()


def _inspect_excel_template(path: Path) -> TemplateLayout:
    workbook = load_workbook(path)
    placeholder_fields: list[str] = []
    label_fields: list[str] = []
    table_targets: list[TemplateTableTarget] = []

    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            values = [normalize_text(cell.value) for cell in row]
            non_empty = [value for value in values if value]
            if not non_empty:
                continue

            for value in non_empty:
                placeholder_fields.extend(_extract_placeholders(value))

        header_row = _find_header_row(sheet)
        if header_row:
            row_index, headers, header_col_map = header_row
            context_rows = _collect_excel_context_rows(sheet, row_index)
            existing_rows = _collect_excel_existing_rows(sheet, row_index + 1)
            target_city = detect_target_city([sheet.title, *context_rows, *existing_rows])
            table_targets.append(
                TemplateTableTarget(
                    template_type="xlsx",
                    identifier=f"{sheet.title}:{row_index}",
                    sheet_name=sheet.title,
                    headers=headers,
                    start_row=row_index + 1,
                    capacity=max(0, sheet.max_row - row_index),
                    context_text="\n".join(context_rows),
                    target_city=target_city,
                    metadata={"header_col_map": header_col_map},
                )
            )
            label_fields.extend(headers)

    mode = "mixed"
    if table_targets and not placeholder_fields:
        mode = "table"
    elif placeholder_fields and not table_targets:
        mode = "placeholder"

    return TemplateLayout(
        template_path=str(path),
        template_type="xlsx",
        mode=mode,
        placeholder_fields=_dedupe(placeholder_fields),
        label_fields=_dedupe(label_fields),
        table_targets=table_targets,
    )


def _inspect_word_template(path: Path) -> TemplateLayout:
    document = Document(path)
    placeholder_fields: list[str] = []
    label_fields: list[str] = []
    table_targets: list[TemplateTableTarget] = []
    placeholder_contexts: dict[str, str] = {}
    body = document.element.body
    table_index = 0
    recent_context: list[str] = []

    for child in body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, document)
            text = normalize_text(paragraph.text)
            if text:
                recent_context.append(text)
                recent_context = recent_context[-6:]
                placeholder_fields.extend(_extract_placeholders(text))
                _collect_placeholder_contexts(text, placeholder_contexts)
            continue

        if child.tag.endswith("}tbl"):
            table = Table(child, document)
            if not table.rows:
                table_index += 1
                continue
            headers = [normalize_text(cell.text) for cell in table.rows[0].cells]
            non_empty_headers = [header for header in headers if header]
            # 保留原始列索引映射：[(原始列号, 表头文本), ...]
            header_col_map = [(ci, h) for ci, h in enumerate(headers) if h]
            if len(non_empty_headers) >= 2:
                context_lines = recent_context[-4:]
                existing_rows = _collect_word_table_texts(table)
                target_city = detect_target_city([*context_lines, *existing_rows, *non_empty_headers])
                # 检测是否为标签式表格（数据行的首列有标签、其余列含 【填写】）
                is_labeled = _is_labeled_table(table)
                table_targets.append(
                    TemplateTableTarget(
                        template_type="docx",
                        identifier=f"table:{table_index}",
                        table_index=table_index,
                        headers=non_empty_headers,
                        start_row=2,
                        capacity=max(0, len(table.rows) - 1),
                        context_text="\n".join(context_lines),
                        target_city=target_city,
                        metadata={"header_col_map": header_col_map},
                    )
                )
                if is_labeled:
                    table_targets[-1].metadata["labeled"] = True
                label_fields.extend(non_empty_headers)
            for row in table.rows:
                for cell in row.cells:
                    placeholder_fields.extend(_extract_placeholders(cell.text))
                    if cell.text.strip():
                        label_fields.append(cell.text.strip())
            table_index += 1
            recent_context = []  # 每张表处理完后重置，避免后续表继承前面表的上下文

    mode = "mixed"
    if table_targets and not placeholder_fields:
        mode = "table"
    elif placeholder_fields and not table_targets:
        mode = "placeholder"

    return TemplateLayout(
        template_path=str(path),
        template_type="docx",
        mode=mode,
        placeholder_fields=_dedupe(placeholder_fields),
        label_fields=_dedupe(label_fields),
        table_targets=table_targets,
        placeholder_contexts=placeholder_contexts,
    )


def _find_header_row(sheet) -> tuple[int, list[str], list[tuple[int, str]]] | None:
    for row in sheet.iter_rows():
        values = [normalize_text(cell.value) for cell in row]
        non_empty = [value for value in values if value]
        if len(non_empty) >= 2:
            if any(value.startswith("{{") and value.endswith("}}") for value in non_empty):
                continue
            # 保留原始列号映射：[(列号, 表头文本), ...]
            header_col_map = [(cell.column, normalize_text(cell.value)) for cell in row if normalize_text(cell.value)]
            return row[0].row, non_empty, header_col_map
    return None


def _collect_excel_context_rows(sheet, header_row_index: int, limit: int = 4) -> list[str]:
    rows: list[str] = []
    for row_no in range(1, header_row_index):
        values = [normalize_text(sheet.cell(row=row_no, column=col).value) for col in range(1, sheet.max_column + 1)]
        text = " | ".join(value for value in values if value)
        if text:
            rows.append(text)
    return rows[-limit:]


def _collect_excel_existing_rows(sheet, start_row: int, limit: int = 4) -> list[str]:
    rows: list[str] = []
    for row_no in range(start_row, min(sheet.max_row, start_row + limit - 1) + 1):
        values = [normalize_text(sheet.cell(row=row_no, column=col).value) for col in range(1, sheet.max_column + 1)]
        text = " | ".join(value for value in values if value)
        if text:
            rows.append(text)
    return rows


def _collect_word_table_texts(table, limit: int = 6) -> list[str]:
    rows: list[str] = []
    for row in table.rows[1:]:
        values = [normalize_text(cell.text) for cell in row.cells]
        text = " | ".join(value for value in values if value)
        if text:
            rows.append(text)
        if len(rows) >= limit:
            break
    return rows


def _is_labeled_table(table) -> bool:
    """判断表格是否为标签式表格（数据行首列有标签，其余列含填空标记或为空）。"""
    if len(table.rows) < 2:
        return False
    labeled_rows = 0
    for row in table.rows[1:]:
        cells = [normalize_text(cell.text) for cell in row.cells]
        if not cells:
            continue
        first_cell = cells[0]
        # 首列有中文文字标签
        has_label = bool(re.search(r"[\u4e00-\u9fff]", first_cell))
        # 其余列含 【xxx】 或 ____/[ ]/（） 或为空
        has_blank = any(
            "【" in c
            or UNDERLINE_BLANK_PATTERN.search(c)
            or BRACKET_BLANK_PATTERN.search(c)
            or CN_PAREN_BLANK_PATTERN.search(c)
            or not c.strip()
            for c in cells[1:]
        )
        if has_label and has_blank:
            labeled_rows += 1
    return labeled_rows >= 1


def _collect_placeholder_contexts(text: str, contexts: dict[str, str]) -> None:
    """为段落中的每个 【xxx】 占位符收集上下文文本。

    上下文文本 = 段落原文去除所有 【xxx】、{{xxx}} 后的纯文本，
    作为后续检索和模糊匹配的语义依据，避免占位符本身的示例/单位说明干扰检索。
    同一占位符文本出现在多个段落时，上下文会用 " | " 串联起来。
    """
    if "\u3010" not in text:
        return
    cn_matches = [m.group(1).strip() for m in CN_PLACEHOLDER_PATTERN.finditer(text)]
    cn_matches = [ph for ph in cn_matches if ph and ph not in _GENERIC_TABLE_MARKERS]
    if not cn_matches:
        return
    pure = CN_PLACEHOLDER_PATTERN.sub("", text)
    pure = PLACEHOLDER_PATTERN.sub("", pure)
    pure = pure.strip()
    if not pure:
        return
    for ph in cn_matches:
        existing = contexts.get(ph, "")
        if not existing:
            contexts[ph] = pure
        elif pure not in existing:
            contexts[ph] = existing + " | " + pure


def _extract_placeholders(text: str) -> list[str]:
    """从文本中提取所有占位符字段名。

    支持的占位符形式：
    - {{field_name}}
    - 【field_description】
    - ____（下划线填空）
    - [ ]（方括号填空）
    - （）（中文圆括号填空）

    对下划线/括号类填空，用前后文生成字段名。
    """
    normalized = normalize_text(text)
    results = [match.group(1).strip() for match in PLACEHOLDER_PATTERN.finditer(normalized)]
    for match in CN_PLACEHOLDER_PATTERN.finditer(normalized):
        hint = match.group(1).strip()
        if hint not in _GENERIC_TABLE_MARKERS:
            results.append(hint)
    # 下划线填空
    for match in UNDERLINE_BLANK_PATTERN.finditer(normalized):
        field_name = _context_field_name(normalized, match.start(), match.end(), "____")
        if field_name:
            results.append(field_name)
    # 方括号填空
    for match in BRACKET_BLANK_PATTERN.finditer(normalized):
        field_name = _context_field_name(normalized, match.start(), match.end(), "[]")
        if field_name:
            results.append(field_name)
    # 中文圆括号填空
    for match in CN_PAREN_BLANK_PATTERN.finditer(normalized):
        field_name = _context_field_name(normalized, match.start(), match.end(), "（）")
        if field_name:
            results.append(field_name)
    return results


def _context_field_name(text: str, start: int, end: int, marker: str) -> str:
    """从占位符前后文提取上下文字段名。

    提取前后文的中文/英文关键短语作为字段名提示，
    去除常见动词和虚词后生成 "前文__后文" 格式。
    """
    before_raw = text[max(0, start - 30):start]
    after_raw = text[end:end + 20]
    # 去除前文尾部标点
    before = re.sub(r'[，,：:；;、。！？\s]+$', '', before_raw)
    # 去除前文尾部常见动词/虚词（不去掉有实际语义的词如"增长"）
    before = re.sub(r'(?:达到|为|是|共|计|约|有|近|至|了|的|其中|分别)$', '', before)
    # 提取前文末尾的中文/英文/数字短语
    before_m = re.search(r'([\u4e00-\u9fffA-Za-z0-9%％]{2,15})$', before)
    b = before_m.group(1) if before_m else ""
    # 去除后文开头标点，提取后文开头文字
    after = re.sub(r'^[，,：:；;、。！？\s]+', '', after_raw)
    after_m = re.search(r'^([\u4e00-\u9fffA-Za-z%％°]{1,10})', after)
    a = after_m.group(1) if after_m else ""
    if not b and not a:
        return ""
    if b and a:
        return f"{b}__{a}"
    return b or a


def _dedupe(items: list[str]) -> list[str]:
    result: list[str] = []
    seen: set[str] = set()
    for item in items:
        key = normalize_field_name(item)
        if not key or key in seen:
            continue
        seen.add(key)
        result.append(item)
    return result


# ─── 通用空白填充 ──────────────────────────────────────────────


def fill_blanks_in_text(text: str, values: dict[str, object]) -> str:
    """在一段文本中查找 ____/[ ]/（） 空白并用最优匹配值替换。

    对每个空白：
      1. 提取前后上下文关键词
      2. 在 values 字典中用多策略评分查找最佳候选
      3. 从右向左替换（避免偏移问题）
    """
    if not values:
        return text

    blanks: list[tuple[int, int]] = []
    for pattern in (UNDERLINE_BLANK_PATTERN, BRACKET_BLANK_PATTERN, CN_PAREN_BLANK_PATTERN):
        for m in pattern.finditer(text):
            blanks.append((m.start(), m.end()))

    if not blanks:
        return text

    blanks.sort(key=lambda x: x[0], reverse=True)
    used_keys: set[str] = set()

    for start, end in blanks:
        before = text[max(0, start - 40):start]
        after = text[end:end + 30]
        best_key = _match_blank_to_value(before, after, values, used_keys)
        if best_key:
            replacement = normalize_text(values[best_key])
            used_keys.add(best_key)
            text = text[:start] + replacement + text[end:]

    return text


def _match_blank_to_value(
    before: str,
    after: str,
    values: dict[str, object],
    used_keys: set[str],
) -> str:
    """针对一个空白位置，在 values 中找到最匹配的值。

    多策略评分：
      1. 上下文 key 精确匹配
      2. 归一化子串匹配
      3. 关键词重叠度评分（近处的词权重更高）
    """
    if not values:
        return ""

    before_clean = re.sub(r'[，,：:；;、。！？\s]+$', '', before)
    after_clean = re.sub(r'^[，,：:；;、。！？\s]+', '', after)

    # 生成 context_key 尝试精确匹配
    stripped_before = re.sub(
        r'(?:达到|为|是|共|计|约|有|近|至|了|的|其中|分别)$', '', before_clean
    )
    before_m = re.search(r'([\u4e00-\u9fffA-Za-z0-9%％]{2,15})$', stripped_before)
    after_m = re.search(r'^([\u4e00-\u9fffA-Za-z%％°]{1,10})', after_clean)
    b = before_m.group(1) if before_m else ""
    a = after_m.group(1) if after_m else ""
    context_key = f"{b}__{a}" if b and a else (b or a)
    if context_key and context_key in values and context_key not in used_keys:
        return context_key

    # 提取上下文关键词
    before_words = re.findall(r'[\u4e00-\u9fff]{2,}|[A-Za-z0-9]+', before_clean)
    after_words = re.findall(r'[\u4e00-\u9fff]{2,}|[A-Za-z0-9]+', after_clean)
    norm_context = normalize_field_name(before_clean + after_clean)

    best_key = ""
    best_score = 0.0

    for key in values:
        if key in used_keys:
            continue
        nk = normalize_field_name(key)
        if not nk:
            continue

        score = 0.0

        # 策略 1: key 原文出现在上下文中
        if key in before_clean or key in after_clean:
            score += 6

        # 策略 2: 归一化 key 是上下文子串（或反向包含）
        if nk in norm_context:
            score += 4
        elif norm_context and norm_context in nk:
            score += 3

        # 策略 3: 前文关键词匹配（近处的词权重更高）
        focus = before_words[-4:]
        for idx, word in enumerate(focus):
            weight = 1.0 + idx * 0.5  # idx 越大 = 越靠近空白 = 权重越高
            nw = normalize_field_name(word)
            if not nw:
                continue
            if nw in nk:
                score += 3 * weight
            elif nk in nw:
                score += 2 * weight
            elif len(nw) >= 2:
                for i in range(len(nw) - 1):
                    if nw[i:i + 2] in nk:
                        score += 0.5 * weight
                        break

        # 策略 4: 后文关键词匹配
        for word in after_words[:2]:
            nw = normalize_field_name(word)
            if nw and nw in nk:
                score += 2

        if score > best_score:
            best_score = score
            best_key = key

    return best_key if best_score >= 3 else ""


def has_blank_marker(text: str) -> bool:
    """检查文本是否包含空白填充标记（____/[ ]/（））。"""
    return bool(
        UNDERLINE_BLANK_PATTERN.search(text)
        or BRACKET_BLANK_PATTERN.search(text)
        or CN_PAREN_BLANK_PATTERN.search(text)
    )


def build_output_path(template_path: str, output_dir: str) -> Path:
    template = Path(template_path)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_root = Path(output_dir)
    output_root.mkdir(parents=True, exist_ok=True)
    return output_root / f"{template.stem}_filled_{timestamp}{template.suffix}"


def build_report_path(output_path: str | Path) -> Path:
    return Path(output_path).with_suffix(".json")


def save_json(data, path: str | Path) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(
        json.dumps(data, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
