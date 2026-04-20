"""
excel_filler.py — Excel 模板填写器

三种填写方式：
  1. 占位符替换：{{field_name}} → value
  2. 标签-值对填写：“字段名”单元格右侧填入值
  3. 表格数据填写：按表头列增写多行记录
"""

from __future__ import annotations

import logging

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

logger = logging.getLogger(__name__)

from app.filler.utils import inspect_template, fill_blanks_in_text
from app.schemas import TemplateLayout, normalize_field_name, normalize_text


def fill_excel(
    template_path: str,
    records: list[dict[str, object]],
    output_path: str,
    layout: TemplateLayout | None = None,
    single_record: dict[str, object] | None = None,
    table_records: dict[str, list[dict[str, object]]] | None = None,
) -> None:
    workbook = load_workbook(template_path)
    layout = layout or inspect_template(template_path)
    single_record = single_record or (records[0] if records else {})
    table_records = table_records or {}

    for sheet in workbook.worksheets:
        _replace_placeholders(sheet, single_record)
        _fill_label_value_pairs(sheet, single_record)
        _replace_blank_patterns(sheet, single_record)

    for target in layout.table_targets:
        if target.template_type != "xlsx":
            continue
        sheet = workbook[target.sheet_name]
        target_records = table_records.get(target.identifier, records)
        header_col_map = target.metadata.get("header_col_map")
        _fill_table(sheet, target.headers, target.start_row, target_records, header_col_map=header_col_map)

    workbook.save(output_path)


def _replace_placeholders(sheet, record: dict[str, object]) -> None:
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            text = normalize_text(cell.value)
            if not text:
                continue
            for key, value in record.items():
                placeholder = "{{" + key + "}}"
                if placeholder in text:
                    cell.value = text.replace(placeholder, normalize_text(value))


def _fill_label_value_pairs(sheet, record: dict[str, object]) -> None:
    normalized_record = {
        normalize_field_name(key): normalize_text(value)
        for key, value in record.items()
    }
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            label = normalize_field_name(normalize_text(cell.value))
            if not label or label not in normalized_record:
                continue
            target_value = normalized_record[label]
            right_cell = sheet.cell(row=cell.row, column=cell.column + 1)
            if isinstance(right_cell, MergedCell):
                continue
            if not normalize_text(right_cell.value):
                right_cell.value = target_value


def _fuzzy_match_record_key(
    header_key: str,
    record_map: dict[str, object],
) -> str:
    """当精确 normalize 匹配失败时，尝试包含匹配。"""
    if not header_key:
        return ""
    best_key = ""
    best_overlap = 0
    for rk, rv in record_map.items():
        if not rk or not rv:
            continue
        shorter, longer = (header_key, rk) if len(header_key) <= len(rk) else (rk, header_key)
        if shorter in longer:
            overlap = len(shorter) / len(longer)
            if overlap >= 0.5 and len(rk) > best_overlap:
                best_key = rk
                best_overlap = len(rk)
    if best_key:
        return normalize_text(record_map[best_key])
    return ""


def _fill_table(
    sheet,
    headers: list[str],
    start_row: int,
    records: list[dict[str, object]],
    header_col_map: list[tuple[int, str]] | None = None,
) -> None:
    if not records:
        return

    # 使用原始列号映射（若有），否则回退到从第 1 列顺序排列
    if header_col_map:
        col_mapping = [(col_no, normalize_field_name(h)) for col_no, h in header_col_map]
    else:
        col_mapping = [(i + 1, normalize_field_name(h)) for i, h in enumerate(headers)]

    column_matched = [False] * len(col_mapping)

    for record_index, record in enumerate(records, start=0):
        row_no = start_row + record_index
        record_map = {
            normalize_field_name(key): value
            for key, value in record.items()
        }
        for map_index, (actual_col, header_key) in enumerate(col_mapping):
            value = record_map.get(header_key, "")
            if not value:
                value = _fuzzy_match_record_key(header_key, record_map)
            if value:
                column_matched[map_index] = True
            target_cell = sheet.cell(row=row_no, column=actual_col)
            if isinstance(target_cell, MergedCell):
                continue
            target_cell.value = normalize_text(value)

    for mi, (matched, (_, hdr_key)) in enumerate(zip(column_matched, col_mapping)):
        if not matched and hdr_key:
            orig_hdr = headers[mi] if mi < len(headers) else hdr_key
            sample_keys = list(records[0].keys())[:8] if records else []
            logger.warning(
                "[Excel填表] 列 '%s' 在所有记录中均未匹配到值，记录字段示例: %s",
                orig_hdr, sample_keys,
            )


def _replace_blank_patterns(sheet, record: dict[str, object]) -> None:
    """替换 Excel 单元格中的 ____/[ ]/（） 空白占位符。"""
    if not record:
        return
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            text = normalize_text(cell.value)
            if not text:
                continue
            updated = fill_blanks_in_text(text, record)
            if updated != text:
                cell.value = updated
